#!/usr/bin/env python3
"""Exports all 82 slides of Project AEOLUS to a high-fidelity PowerPoint (.pptx) presentation.

Optimized for 16:9 widescreen presentation, tablet viewing, and native Google Slides import.
"""

import os
import sys
import shutil
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

# Add project root to sys.path
BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)

from deck_builder.part0_intro import get_part0_slides
from deck_builder.part1_governing import get_part1_slides
from deck_builder.part2_microphysics import get_part2_slides
from deck_builder.part3_kinematics import get_part3_slides
from deck_builder.part4_cfd_verification import get_part4_slides
from deck_builder.part5_code_ledger import get_part5_slides
from deck_builder.part6_hardware_prototype import get_part6_slides

# Color Palette
PALETTE = {
    "BG_DARK": RGBColor(11, 17, 32),       # Deep Obsidian Slate (#0B1120)
    "CARD_BG": RGBColor(24, 34, 53),       # Elevated Card Slate (#182235)
    "CARD_BORDER": RGBColor(51, 65, 85),   # Muted Slate Border (#334155)
    "CODE_BG": RGBColor(7, 11, 20),        # Monospace Deep Dark (#070B14)
    "ACCENT_CYAN": RGBColor(56, 189, 248), # #38BDF8
    "ACCENT_GREEN": RGBColor(52, 211, 153),# #34D399
    "ACCENT_AMBER": RGBColor(245, 158, 11),# #F59E0B
    "ACCENT_CORAL": RGBColor(248, 113, 113),# #F87171
    "TEXT_WHITE": RGBColor(255, 255, 255),
    "TEXT_BODY": RGBColor(226, 232, 240),  # #E2E8F0
    "TEXT_MUTED": RGBColor(148, 163, 184)  # #94A3B8
}

def hex_to_rgb(hex_str):
    if not hex_str:
        return PALETTE["ACCENT_CYAN"]
    hex_str = hex_str.lstrip("#")
    try:
        return RGBColor(int(hex_str[0:2], 16), int(hex_str[2:4], 16), int(hex_str[4:6], 16))
    except Exception:
        return PALETTE["ACCENT_CYAN"]

def create_base_slide(prs):
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = PALETTE["BG_DARK"]
    bg.line.fill.background()
    return slide

def apply_header(slide, slide_num, total_slides, part_label, title_text, subtitle_text):
    # Top rule accent line
    top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.06))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = PALETTE["ACCENT_CYAN"]
    top_bar.line.fill.background()

    # Part & Slide Number Tag Box
    tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.28), Inches(11.733), Inches(0.32))
    tf_tag = tag_box.text_frame
    tf_tag.word_wrap = True
    tf_tag.margin_left = tf_tag.margin_right = tf_tag.margin_top = tf_tag.margin_bottom = 0
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = f"SLIDE {slide_num} OF {total_slides}  •  {part_label.upper()}"
    p_tag.font.size = Pt(10)
    p_tag.font.bold = True
    p_tag.font.color.rgb = PALETTE["ACCENT_CYAN"]

    # Title Box
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.58), Inches(11.733), Inches(0.55))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    tf_title.margin_left = tf_title.margin_right = tf_title.margin_top = tf_title.margin_bottom = 0
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.size = Pt(21)
    p_title.font.bold = True
    p_title.font.color.rgb = PALETTE["TEXT_WHITE"]

    # Subtitle Box
    if subtitle_text:
        sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.15), Inches(11.733), Inches(0.35))
        tf_sub = sub_box.text_frame
        tf_sub.word_wrap = True
        tf_sub.margin_left = tf_sub.margin_right = tf_sub.margin_top = tf_sub.margin_bottom = 0
        p_sub = tf_sub.paragraphs[0]
        p_sub.text = subtitle_text
        p_sub.font.size = Pt(11)
        p_sub.font.italic = True
        p_sub.font.color.rgb = PALETTE["TEXT_MUTED"]

def insert_card(slide, x, y, w, h, bg_color=None, border_color=None):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color or PALETTE["CARD_BG"]
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = Pt(1.2)
    else:
        card.line.color.rgb = PALETTE["CARD_BORDER"]
        card.line.width = Pt(1)
    return card

def add_bullets_to_card(card, title, bullets, accent_rgb):
    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.3)
    tf.margin_right = Inches(0.3)
    tf.margin_top = Inches(0.28)
    tf.margin_bottom = Inches(0.28)

    p_title = tf.paragraphs[0]
    p_title.text = title
    p_title.font.size = Pt(13)
    p_title.font.bold = True
    p_title.font.color.rgb = accent_rgb
    p_title.space_after = Pt(10)

    for b in bullets:
        raw = b.strip()
        if not raw:
            continue
        p = tf.add_paragraph()
        if raw.startswith("•"):
            text = raw.lstrip("•").strip()
            p.text = "•  " + text
            p.font.size = Pt(11)
            p.font.bold = True
            p.font.color.rgb = PALETTE["TEXT_WHITE"]
            p.space_before = Pt(6)
            p.space_after = Pt(2)
        elif raw.startswith("-"):
            text = raw.lstrip("-").strip()
            p.text = "     ▪  " + text
            p.font.size = Pt(10.2)
            p.font.bold = False
            p.font.color.rgb = PALETTE["TEXT_BODY"]
            p.space_after = Pt(3)
        else:
            p.text = "•  " + raw
            p.font.size = Pt(10.5)
            p.font.color.rgb = PALETTE["TEXT_BODY"]
            p.space_after = Pt(3)

def insert_code_box(slide, x, y, w, h, header_title, code_string, accent_rgb):
    box = insert_card(slide, x, y, w, h, PALETTE["CODE_BG"], accent_rgb)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.25)
    tf.margin_right = Inches(0.25)
    tf.margin_top = Inches(0.25)
    tf.margin_bottom = Inches(0.25)

    if header_title:
        p_hdr = tf.paragraphs[0]
        p_hdr.text = header_title
        p_hdr.font.size = Pt(11)
        p_hdr.font.bold = True
        p_hdr.font.color.rgb = accent_rgb
        p_hdr.space_after = Pt(8)
        p_code = tf.add_paragraph()
    else:
        p_code = tf.paragraphs[0]

    p_code.text = code_string
    p_code.font.name = "Courier New"
    p_code.font.size = Pt(9.2)
    p_code.font.color.rgb = PALETTE["TEXT_BODY"]
    return box

def set_speaker_notes(slide, notes_content):
    if notes_content:
        notes_slide = slide.notes_slide
        tf = notes_slide.notes_text_frame
        tf.text = notes_content

def render_title_slide(slide, s, slide_num, total):
    # Top bar
    top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.08))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = PALETTE["ACCENT_CYAN"]
    top_bar.line.fill.background()

    # Category Badge
    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.4), Inches(4.5), Inches(0.4))
    badge.fill.solid()
    badge.fill.fore_color.rgb = PALETTE["CARD_BG"]
    badge.line.color.rgb = PALETTE["ACCENT_CYAN"]
    badge.line.width = Pt(1)
    tf_b = badge.text_frame
    tf_b.margin_top = Inches(0.05)
    p_b = tf_b.paragraphs[0]
    p_b.text = "COMPUTATIONAL FLUID DYNAMICS & HARDWARE DEFENSE"
    p_b.font.size = Pt(9.5)
    p_b.font.bold = True
    p_b.font.color.rgb = PALETTE["ACCENT_CYAN"]
    p_b.alignment = PP_ALIGN.CENTER

    # Main Title
    tbox = slide.shapes.add_textbox(Inches(0.8), Inches(0.9), Inches(11.733), Inches(0.8))
    tf_t = tbox.text_frame
    tf_t.word_wrap = True
    p_t = tf_t.paragraphs[0]
    p_t.text = s.get("title", "PROJECT AEOLUS")
    p_t.font.size = Pt(32)
    p_t.font.bold = True
    p_t.font.color.rgb = PALETTE["TEXT_WHITE"]

    # Subtitle
    sbox = slide.shapes.add_textbox(Inches(0.8), Inches(1.65), Inches(11.733), Inches(0.45))
    tf_s = sbox.text_frame
    tf_s.word_wrap = True
    p_s = tf_s.paragraphs[0]
    p_s.text = s.get("subtitle", "")
    p_s.font.size = Pt(12.5)
    p_s.font.italic = True
    p_s.font.color.rgb = PALETTE["TEXT_MUTED"]

    # Stats Row (4 Hero Cards)
    stats = s.get("stats", [])
    card_w = 2.75
    gap = 0.244
    for i, stat in enumerate(stats):
        cx = 0.8 + i * (card_w + gap)
        scard = insert_card(slide, cx, 2.25, card_w, 1.6, PALETTE["CARD_BG"], hex_to_rgb(stat.get("col")))
        tf_sc = scard.text_frame
        tf_sc.word_wrap = True
        tf_sc.margin_left = Inches(0.2)
        tf_sc.margin_right = Inches(0.2)
        tf_sc.margin_top = Inches(0.25)
        tf_sc.margin_bottom = Inches(0.2)

        p_v = tf_sc.paragraphs[0]
        p_v.text = stat.get("val", "")
        p_v.font.size = Pt(24)
        p_v.font.bold = True
        p_v.font.color.rgb = hex_to_rgb(stat.get("col"))
        p_v.alignment = PP_ALIGN.CENTER

        p_l = tf_sc.add_paragraph()
        p_l.text = stat.get("lbl", "")
        p_l.font.size = Pt(10)
        p_l.font.color.rgb = PALETTE["TEXT_MUTED"]
        p_l.alignment = PP_ALIGN.CENTER
        p_l.space_before = Pt(4)

    # Executive Statement Card
    exec_card = insert_card(slide, 0.8, 4.1, 11.733, 2.85, PALETTE["CARD_BG"], PALETTE["CARD_BORDER"])
    tf_ex = exec_card.text_frame
    tf_ex.word_wrap = True
    tf_ex.margin_left = Inches(0.35)
    tf_ex.margin_right = Inches(0.35)
    tf_ex.margin_top = Inches(0.3)
    tf_ex.margin_bottom = Inches(0.3)

    p_eh = tf_ex.paragraphs[0]
    p_eh.text = "EXECUTIVE TECHNICAL BRIEFING STATEMENT"
    p_eh.font.size = Pt(13)
    p_eh.font.bold = True
    p_eh.font.color.rgb = PALETTE["ACCENT_CYAN"]
    p_eh.space_after = Pt(10)

    p_eb = tf_ex.add_paragraph()
    p_eb.text = s.get("summary", "")
    p_eb.font.size = Pt(11.5)
    p_eb.font.color.rgb = PALETTE["TEXT_BODY"]

    set_speaker_notes(slide, s.get("notes", ""))

def render_three_cards(slide, s, slide_num, total):
    apply_header(slide, slide_num, total, s.get("part", ""), s.get("title", ""), s.get("subtitle", ""))
    cards = s.get("cards", [])
    card_w = 3.73
    gap = 0.27
    for i, c in enumerate(cards):
        cx = 0.8 + i * (card_w + gap)
        accent_rgb = hex_to_rgb(c.get("col"))
        card = insert_card(slide, cx, 1.65, card_w, 5.35, PALETTE["CARD_BG"], accent_rgb)
        add_bullets_to_card(card, c.get("title", ""), c.get("bullets", []), accent_rgb)
    set_speaker_notes(slide, s.get("notes", ""))

def render_split_cards(slide, s, slide_num, total):
    apply_header(slide, slide_num, total, s.get("part", ""), s.get("title", ""), s.get("subtitle", ""))
    card1 = s.get("card1", {})
    card2 = s.get("card2", {})
    card_w = 5.72
    gap = 0.29

    col1 = hex_to_rgb(card1.get("col"))
    c1 = insert_card(slide, 0.8, 1.65, card_w, 5.35, PALETTE["CARD_BG"], col1)
    add_bullets_to_card(c1, card1.get("title", ""), card1.get("bullets", []), col1)

    col2 = hex_to_rgb(card2.get("col"))
    c2 = insert_card(slide, 0.8 + card_w + gap, 1.65, card_w, 5.35, PALETTE["CARD_BG"], col2)
    add_bullets_to_card(c2, card2.get("title", ""), card2.get("bullets", []), col2)

    set_speaker_notes(slide, s.get("notes", ""))

def render_split_code(slide, s, slide_num, total):
    apply_header(slide, slide_num, total, s.get("part", ""), s.get("title", ""), s.get("subtitle", ""))
    accent_rgb = hex_to_rgb(s.get("accent"))
    card_w = 5.72
    gap = 0.29

    insert_code_box(slide, 0.8, 1.65, card_w, 5.35, s.get("codeHeader", ""), s.get("codeText", ""), accent_rgb)

    card_data = s.get("card", {})
    card_rgb = hex_to_rgb(card_data.get("col") or s.get("accent"))
    c = insert_card(slide, 0.8 + card_w + gap, 1.65, card_w, 5.35, PALETTE["CARD_BG"], card_rgb)
    add_bullets_to_card(c, card_data.get("title", ""), card_data.get("bullets", []), card_rgb)

    set_speaker_notes(slide, s.get("notes", ""))

def render_full_code(slide, s, slide_num, total):
    apply_header(slide, slide_num, total, s.get("part", ""), s.get("title", ""), s.get("subtitle", ""))
    accent_rgb = hex_to_rgb(s.get("accent"))

    insert_code_box(slide, 0.8, 1.65, 11.733, 3.55, s.get("codeHeader", ""), s.get("codeText", ""), accent_rgb)

    bottom_data = s.get("bottomCard", {})
    b_rgb = hex_to_rgb(bottom_data.get("col") or s.get("accent"))
    c = insert_card(slide, 0.8, 5.35, 11.733, 1.65, PALETTE["CARD_BG"], b_rgb)
    add_bullets_to_card(c, bottom_data.get("title", ""), bottom_data.get("bullets", []), b_rgb)

    set_speaker_notes(slide, s.get("notes", ""))

def render_table_slide(slide, s, slide_num, total):
    apply_header(slide, slide_num, total, s.get("part", ""), s.get("title", ""), s.get("subtitle", ""))
    headers = s.get("headers", [])
    rows = s.get("rows", [])
    col_widths = s.get("colWidths", [])

    num_rows = len(rows) + 1
    num_cols = len(headers)

    table_shape = slide.shapes.add_table(num_rows, num_cols, Inches(0.8), Inches(1.65), Inches(11.733), Inches(5.2))
    table = table_shape.table

    if col_widths and len(col_widths) == num_cols:
        total_cw = sum(col_widths)
        for c in range(num_cols):
            table.columns[c].width = int(Inches(11.733) * (col_widths[c] / total_cw))

    # Headers
    for c in range(num_cols):
        cell = table.cell(0, c)
        cell.text = headers[c]
        cell.fill.solid()
        cell.fill.fore_color.rgb = PALETTE["CARD_BG"]
        p = cell.text_frame.paragraphs[0]
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = PALETTE["ACCENT_CYAN"]

    # Rows
    for r in range(len(rows)):
        is_highlight = (r == len(rows) - 1)
        for c in range(num_cols):
            cell = table.cell(r + 1, c)
            cell.text = rows[r][c]
            cell.fill.solid()
            if is_highlight:
                cell.fill.fore_color.rgb = RGBColor(6, 78, 59) # #064E3B
            elif r % 2 == 0:
                cell.fill.fore_color.rgb = PALETTE["BG_DARK"]
            else:
                cell.fill.fore_color.rgb = PALETTE["CARD_BG"]

            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(10.5)
            p.font.bold = is_highlight
            p.font.color.rgb = PALETTE["ACCENT_GREEN"] if is_highlight else PALETTE["TEXT_BODY"]

    set_speaker_notes(slide, s.get("notes", ""))

def compile_all_slides():
    slides = []
    slides.extend(get_part0_slides())
    slides.extend(get_part1_slides())
    slides.extend(get_part2_slides())
    slides.extend(get_part3_slides())
    slides.extend(get_part4_slides())
    slides.extend(get_part5_slides())
    slides.extend(get_part6_slides())
    return slides

def export_to_pptx(output_pptx_path):
    print("Compiling all 82 slides for PowerPoint (.pptx)...")
    slides = compile_all_slides()
    total = len(slides)
    print(f"Total slides to generate: {total}")

    prs = pptx.Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.500)

    for i, s in enumerate(slides):
        slide_num = i + 1
        stype = s.get("type", "split_cards")
        slide = create_base_slide(prs)

        if stype == "title":
            render_title_slide(slide, s, slide_num, total)
        elif stype == "three_cards":
            render_three_cards(slide, s, slide_num, total)
        elif stype == "split_cards":
            render_split_cards(slide, s, slide_num, total)
        elif stype == "split_code":
            render_split_code(slide, s, slide_num, total)
        elif stype == "full_code":
            render_full_code(slide, s, slide_num, total)
        elif stype == "table":
            render_table_slide(slide, s, slide_num, total)
        else:
            render_split_cards(slide, s, slide_num, total)

        if slide_num % 10 == 0 or slide_num == total:
            print(f"  Processed {slide_num} / {total} slides...")

    prs.save(output_pptx_path)
    file_size_kb = os.path.getsize(output_pptx_path) / 1024
    print(f"Successfully exported PPTX to {output_pptx_path} ({file_size_kb:.1f} KB)")

    # Copy to tablet Downloads folder if available
    download_dir = "/storage/emulated/0/Download"
    if os.path.isdir(download_dir):
        dest = os.path.join(download_dir, os.path.basename(output_pptx_path))
        shutil.copy2(output_pptx_path, dest)
        print(f"Copied PPTX presentation directly to Android tablet storage: {dest}")

if __name__ == "__main__":
    local_output = os.path.join(BASE_DIR, "Project_AEOLUS_Master_Technical_Briefing.pptx")
    export_to_pptx(local_output)
